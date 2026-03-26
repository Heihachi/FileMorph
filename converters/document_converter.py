# FileMorph/converters/document_converter.py
from pathlib import Path

from docx import Document
from docx2pdf import convert as docx2pdf_convert
from pdf2docx import Converter as PdfToDocxConverter
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from converters.base import BaseConverter


class DocumentConverter(BaseConverter):
    @property
    def supported_input_formats(self) -> list[str]:
        return ["docx", "txt", "pdf"]

    @property
    def supported_output_formats(self) -> list[str]:
        return ["txt", "docx", "pdf"]

    def convert(self, src: Path, dst: Path, options: dict) -> None:
        src_ext = src.suffix.lower().lstrip(".")
        dst_ext = dst.suffix.lower().lstrip(".")

        if src_ext == "txt" and dst_ext == "docx":
            self._txt_to_docx(src, dst)
            return

        if src_ext == "txt" and dst_ext == "pdf":
            self._txt_to_pdf(src, dst)
            return

        if src_ext == "docx" and dst_ext == "txt":
            self._docx_to_txt(src, dst)
            return

        if src_ext == "pdf" and dst_ext == "txt":
            self._pdf_to_txt(src, dst)
            return

        if src_ext == "docx" and dst_ext == "pdf":
            self._docx_to_pdf(src, dst)
            return

        if src_ext == "pdf" and dst_ext == "docx":
            self._pdf_to_docx(src, dst)
            return

        raise ValueError(f"Конвертация {src_ext} -> {dst_ext} пока не поддерживается.")

    def _txt_to_docx(self, src: Path, dst: Path) -> None:
        text = src.read_text(encoding="utf-8", errors="replace")

        document = Document()
        for line in text.splitlines():
            document.add_paragraph(line)

        document.save(dst)

    def _txt_to_pdf(self, src: Path, dst: Path) -> None:
        text = src.read_text(encoding="utf-8", errors="replace")

        pdf = canvas.Canvas(str(dst), pagesize=A4)
        width, height = A4

        left_margin = 40
        top_margin = height - 40
        line_height = 16

        font_name = "Helvetica"
        font_size = 11

        pdf.setFont(font_name, font_size)
        y = top_margin

        for raw_line in text.splitlines():
            wrapped_lines = self._wrap_text(
                raw_line,
                max_width=width - left_margin * 2,
                font_name=font_name,
                font_size=font_size
            )

            if not wrapped_lines:
                wrapped_lines = [""]

            for line in wrapped_lines:
                if y < 40:
                    pdf.showPage()
                    pdf.setFont(font_name, font_size)
                    y = top_margin

                pdf.drawString(left_margin, y, line)
                y -= line_height

        pdf.save()

    def _docx_to_txt(self, src: Path, dst: Path) -> None:
        document = Document(src)
        lines = [paragraph.text for paragraph in document.paragraphs]
        text = "\n".join(lines)
        dst.write_text(text, encoding="utf-8")

    def _pdf_to_txt(self, src: Path, dst: Path) -> None:
        reader = PdfReader(str(src))
        parts: list[str] = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            parts.append(page_text)

        text = "\n\n".join(parts)
        dst.write_text(text, encoding="utf-8")

    def _docx_to_pdf(self, src: Path, dst: Path) -> None:

        try:
            docx2pdf_convert(str(src), str(dst))
        except Exception as e:
            raise RuntimeError(
                "DOCX -> PDF не удалось выполнить.\n\n"
                "Чаще всего причина в том, что:\n"
                "1. Microsoft Word не установлен\n"
                "2. Word не может открыть файл\n"
                "3. файл уже открыт в другой программе\n\n"
                f"Детали: {e}"
            ) from e

    def _pdf_to_docx(self, src: Path, dst: Path) -> None:
        converter = PdfToDocxConverter(str(src))
        try:
            converter.convert(str(dst), start=0, end=None)
        except Exception as e:
            raise RuntimeError(
                "PDF -> DOCX не удалось выполнить.\n\n"
                "На сложных PDF (сканы, таблицы, колонки, нестандартная вёрстка) "
                "результат может быть плохим или конвертация может завершиться ошибкой.\n\n"
                f"Детали: {e}"
            ) from e
        finally:
            converter.close()

    def _wrap_text(
        self,
        text: str,
        max_width: float,
        font_name: str,
        font_size: int
    ) -> list[str]:
        if not text:
            return [""]

        words = text.split()
        if not words:
            return [""]

        lines = []
        current_line = words[0]

        for word in words[1:]:
            test_line = f"{current_line} {word}"
            test_width = stringWidth(test_line, font_name, font_size)

            if test_width <= max_width:
                current_line = test_line
            else:
                lines.append(current_line)
                current_line = word

        lines.append(current_line)
        return lines